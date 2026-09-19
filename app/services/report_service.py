"""Reporting, Document Generation & Export Services for Student Result Analysis System.

Provides institutional-grade PDF report cards, comprehensive multi-column Excel
Tabulation Registers (TR Sheets), and printable Executive Class Performance Summaries.
"""

from io import BytesIO
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

from sqlalchemy.orm import joinedload
from app import db
from app.models.academic import Class, Subject, ClassSubject, AcademicYear
from app.models.student import Student
from app.models.result import Marks, ExamTerm
from app.services.analytics_service import AnalyticsService


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and stamp total page count on PDFs."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))

        # Footer divider line
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(40, 40, self._pagesize[0] - 40, 40)

        # Left: Verification note
        self.drawString(40, 28, "Authentic computerized academic record | Student Result Analysis System (IGNOU)")

        # Right: Page count
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(self._pagesize[0] - 40, 28, page_str)
        self.restoreState()


class ReportService:
    """Core reporting service generating institutional PDFs and Excel registers."""

    # --------------------------------------------------------------------------
    # Milestone 6.1: Individual Student Report Card Data & PDF
    # --------------------------------------------------------------------------

    @classmethod
    def get_student_report_card_data(cls, student_id: int, term_id: int) -> Optional[Dict[str, Any]]:
        """
        Aggregate complete demographic, curriculum, and evaluation data for a student's report card.
        """
        student = db.session.get(Student, student_id)
        term = db.session.get(ExamTerm, term_id)
        if not student or not term:
            return None

        target_class = student.enrolled_class
        academic_year = target_class.academic_year if target_class else None

        # Fetch all class subject mappings for the student's class
        mappings = (
            ClassSubject.query.filter_by(class_id=student.class_id)
            .join(Subject)
            .order_by(Subject.subject_code.asc())
            .all()
        )

        # Fetch existing marks records for this student and term
        marks_records = {
            m.class_subject_id: m
            for m in Marks.query.filter_by(student_id=student_id, term_id=term_id).all()
        }

        subject_rows = []
        total_max_marks = 0.0
        total_obtained_marks = 0.0
        total_credits = 0
        weighted_gp_sum = 0.0
        has_fail = False
        all_absent = True

        for cs in mappings:
            sub = cs.subject_ref
            mark = marks_records.get(cs.class_subject_id)
            credits = sub.credits if sub else 4
            max_in = cs.max_internal_marks
            max_ex = cs.max_external_marks
            max_tot = cs.max_total_marks
            pass_mark = cs.pass_marks

            total_max_marks += max_tot
            total_credits += credits

            if mark:
                if not mark.is_absent:
                    all_absent = False
                in_m = float(mark.internal_marks or 0.0)
                ex_m = float(mark.external_marks or 0.0)
                tot_m = float(mark.total_marks or 0.0)
                pct = float(mark.percentage or 0.0)
                grade = mark.grade or 'F'
                gp = float(mark.grade_point or 0.0)
                is_passed = bool(mark.is_passed)
                is_absent = bool(mark.is_absent)

                total_obtained_marks += tot_m
                weighted_gp_sum += (gp * credits)

                if not is_passed:
                    has_fail = True

                status = 'ABSENT' if is_absent else ('PASS' if is_passed else 'FAIL')

                subject_rows.append({
                    'subject_code': sub.subject_code if sub else 'N/A',
                    'subject_name': sub.subject_name if sub else 'Course Subject',
                    'credits': credits,
                    'max_internal': max_in,
                    'internal_marks': '-' if is_absent else in_m,
                    'max_external': max_ex,
                    'external_marks': '-' if is_absent else ex_m,
                    'max_total': max_tot,
                    'total_marks': '-' if is_absent else tot_m,
                    'pass_marks': pass_mark,
                    'percentage': '-' if is_absent else pct,
                    'grade': 'AB' if is_absent else grade,
                    'grade_point': '-' if is_absent else gp,
                    'is_passed': is_passed,
                    'is_absent': is_absent,
                    'status': status
                })
            else:
                has_fail = True
                subject_rows.append({
                    'subject_code': sub.subject_code if sub else 'N/A',
                    'subject_name': sub.subject_name if sub else 'Course Subject',
                    'credits': credits,
                    'max_internal': max_in,
                    'internal_marks': 'N/A',
                    'max_external': max_ex,
                    'external_marks': 'N/A',
                    'max_total': max_tot,
                    'total_marks': 'N/A',
                    'pass_marks': pass_mark,
                    'percentage': 'N/A',
                    'grade': 'N/A',
                    'grade_point': 'N/A',
                    'is_passed': False,
                    'is_absent': False,
                    'status': 'PENDING'
                })

        overall_percentage = round((total_obtained_marks / total_max_marks * 100), 2) if total_max_marks > 0 else 0.0
        cgpa = round((weighted_gp_sum / total_credits), 2) if total_credits > 0 else 0.0

        if not subject_rows:
            overall_result = "NO COURSES ENROLLED"
        elif all_absent and marks_records:
            overall_result = "ABSENT"
        elif has_fail:
            overall_result = "REAPPEAR / FAILED"
        else:
            if overall_percentage >= 75.0:
                overall_result = "PASSED WITH DISTINCTION"
            elif overall_percentage >= 60.0:
                overall_result = "PASSED IN FIRST DIVISION"
            elif overall_percentage >= 50.0:
                overall_result = "PASSED IN SECOND DIVISION"
            else:
                overall_result = "PASSED"

        # Compute class rank
        class_rank = None
        if target_class and total_obtained_marks > 0:
            df = AnalyticsService.get_class_term_dataframe(target_class.class_id, term_id)
            if not df.empty:
                agg = df.groupby('student_id')['total_marks'].sum().sort_values(ascending=False).reset_index()
                ranks = agg[agg['student_id'] == student_id].index
                if len(ranks) > 0:
                    class_rank = int(ranks[0]) + 1

        return {
            'student_id': student.student_id,
            'enrollment_no': student.enrollment_no,
            'full_name': student.full_name,
            'email': student.email or 'N/A',
            'class_id': target_class.class_id if target_class else None,
            'class_name': target_class.class_name if target_class else 'N/A',
            'section': target_class.section if target_class else 'N/A',
            'display_class': target_class.display_name if target_class else 'N/A',
            'academic_year': academic_year.year_label if academic_year else 'N/A',
            'term_id': term.term_id,
            'term_name': term.term_name,
            'is_locked': term.is_locked,
            'issue_date': datetime.now(timezone.utc).strftime("%d-%b-%Y"),
            'subjects': subject_rows,
            'total_max_marks': total_max_marks,
            'total_obtained_marks': total_obtained_marks,
            'overall_percentage': overall_percentage,
            'total_credits': total_credits,
            'cgpa': cgpa,
            'overall_result': overall_result,
            'has_fail': has_fail,
            'class_rank': class_rank
        }

    @classmethod
    def generate_student_report_card_pdf(cls, student_id: int, term_id: int) -> Optional[bytes]:
        """
        Generate high-resolution institutional-grade PDF report card via ReportLab Platypus.
        """
        data = cls.get_student_report_card_data(student_id, term_id)
        if not data:
            return None

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'UnivTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=15,
            leading=18,
            alignment=1,  # Center
            textColor=colors.HexColor('#0f172a')
        )
        sub_title_style = ParagraphStyle(
            'UnivSubTitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.HexColor('#475569')
        )
        badge_style = ParagraphStyle(
            'CardBadge',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            alignment=1,
            textColor=colors.HexColor('#1e3a8a')
        )
        meta_label = ParagraphStyle(
            'MetaLabel',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#334155')
        )
        meta_val = ParagraphStyle(
            'MetaVal',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#0f172a')
        )
        table_hdr = ParagraphStyle(
            'TblHdr',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            alignment=1,
            textColor=colors.white
        )
        table_cell = ParagraphStyle(
            'TblCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            alignment=0,
            textColor=colors.HexColor('#1e293b')
        )
        table_cell_center = ParagraphStyle(
            'TblCellCenter',
            parent=table_cell,
            alignment=1
        )
        table_cell_bold = ParagraphStyle(
            'TblCellBold',
            parent=table_cell,
            fontName='Helvetica-Bold'
        )

        story = []

        # 1. Institutional Header Banner
        story.append(Paragraph("INDIRA GANDHI NATIONAL OPEN UNIVERSITY", title_style))
        story.append(Spacer(1, 2))
        story.append(Paragraph("Maidan Garhi, New Delhi - 110068 | School of Computer and Information Sciences", sub_title_style))
        story.append(Spacer(1, 6))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1e3a8a'), spaceBefore=2, spaceAfter=8))
        story.append(Paragraph("STUDENT GRADE CARD / OFFICIAL STATEMENT OF MARKS", badge_style))
        story.append(Spacer(1, 10))

        # 2. Student & Cohort Metadata Box
        meta_data = [
            [
                Paragraph("Enrollment No:", meta_label), Paragraph(f"<b>{data['enrollment_no']}</b>", meta_val),
                Paragraph("Academic Year:", meta_label), Paragraph(data['academic_year'], meta_val),
            ],
            [
                Paragraph("Student Name:", meta_label), Paragraph(f"<b>{data['full_name']}</b>", meta_val),
                Paragraph("Examination Term:", meta_label), Paragraph(data['term_name'], meta_val),
            ],
            [
                Paragraph("Program / Class:", meta_label), Paragraph(data['display_class'], meta_val),
                Paragraph("Date of Issue:", meta_label), Paragraph(data['issue_date'], meta_val),
            ]
        ]
        meta_table = Table(meta_data, colWidths=[90, 175, 105, 150])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#f1f5f9')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 12))

        # 3. Evaluation Course Breakdown Table
        table_headers = [
            Paragraph("Course Code", table_hdr),
            Paragraph("Course Title", table_hdr),
            Paragraph("Cr", table_hdr),
            Paragraph("Int<br/>(Max)", table_hdr),
            Paragraph("Int<br/>Scored", table_hdr),
            Paragraph("Ext<br/>(Max)", table_hdr),
            Paragraph("Ext<br/>Scored", table_hdr),
            Paragraph("Total<br/>(Max)", table_hdr),
            Paragraph("Total<br/>Scored", table_hdr),
            Paragraph("Grd", table_hdr),
            Paragraph("GP", table_hdr),
            Paragraph("Result", table_hdr),
        ]

        table_rows = [table_headers]

        for s in data['subjects']:
            status_color = colors.HexColor('#15803d') if s['status'] == 'PASS' else colors.HexColor('#b91c1c')
            status_p = Paragraph(f"<font color='{status_color.hexval()}'><b>{s['status']}</b></font>", table_cell_center)

            table_rows.append([
                Paragraph(s['subject_code'], table_cell_bold),
                Paragraph(s['subject_name'], table_cell),
                Paragraph(str(s['credits']), table_cell_center),
                Paragraph(str(int(s['max_internal'])), table_cell_center),
                Paragraph(str(s['internal_marks']), table_cell_center),
                Paragraph(str(int(s['max_external'])), table_cell_center),
                Paragraph(str(s['external_marks']), table_cell_center),
                Paragraph(str(int(s['max_total'])), table_cell_center),
                Paragraph(str(s['total_marks']), table_cell_center),
                Paragraph(f"<b>{s['grade']}</b>", table_cell_center),
                Paragraph(str(s['grade_point']), table_cell_center),
                status_p
            ])

        # Summary Row
        table_rows.append([
            Paragraph("<b>GRAND TOTAL</b>", table_cell_bold),
            Paragraph(f"<b>{len(data['subjects'])} Courses Evaluated</b>", table_cell),
            Paragraph(f"<b>{data['total_credits']}</b>", table_cell_center),
            Paragraph("-", table_cell_center),
            Paragraph("-", table_cell_center),
            Paragraph("-", table_cell_center),
            Paragraph("-", table_cell_center),
            Paragraph(f"<b>{int(data['total_max_marks'])}</b>", table_cell_center),
            Paragraph(f"<b>{data['total_obtained_marks']:.1f}</b>", table_cell_center),
            Paragraph("-", table_cell_center),
            Paragraph("-", table_cell_center),
            Paragraph(f"<b>{data['overall_result'].split()[0]}</b>", table_cell_center)
        ])

        col_widths = [55, 140, 24, 34, 34, 34, 34, 36, 36, 26, 26, 40]
        marks_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        marks_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#1e3a8a')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f8fafc')]),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e2e8f0')),
        ]))
        story.append(marks_table)
        story.append(Spacer(1, 14))

        # 4. Final Performance Summary Box
        res_color = '#15803d' if not data['has_fail'] else '#b91c1c'
        summary_data = [
            [
                Paragraph("<b>Total Marks Secured:</b>", meta_label),
                Paragraph(f"{data['total_obtained_marks']:.1f} / {data['total_max_marks']:.1f}", meta_val),
                Paragraph("<b>Overall Percentage:</b>", meta_label),
                Paragraph(f"<b>{data['overall_percentage']:.2f}%</b>", meta_val),
            ],
            [
                Paragraph("<b>Cumulative GPA (SGPA):</b>", meta_label),
                Paragraph(f"<b>{data['cgpa']:.2f} / 10.00</b>", meta_val),
                Paragraph("<b>Class Merit Rank:</b>", meta_label),
                Paragraph(f"Rank {data['class_rank']}" if data['class_rank'] else "N/A", meta_val),
            ],
            [
                Paragraph("<b>Final Academic Result:</b>", meta_label),
                Paragraph(f"<font color='{res_color}'><b>{data['overall_result']}</b></font>", meta_val),
                Paragraph("<b>Grading Standard:</b>", meta_label),
                Paragraph("IGNOU Standard 10-Point Scale", meta_val),
            ]
        ]
        sum_table = Table(summary_data, colWidths=[130, 135, 125, 130])
        sum_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f1f5f9')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#94a3b8')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(sum_table)
        story.append(Spacer(1, 24))

        # 5. Signatures & Verification Block
        sig_data = [
            [
                Paragraph("__________________________<br/><b>Verified by (Evaluator)</b><br/><font size=7 color='#64748b'>Examination In-Charge</font>", ParagraphStyle('Sig1', parent=styles['Normal'], alignment=0)),
                Paragraph("__________________________<br/><b>Registrar / Director</b><br/><font size=7 color='#64748b'>Controller of Examinations</font>", ParagraphStyle('Sig2', parent=styles['Normal'], alignment=2)),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[260, 260])
        sig_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'BOTTOM'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(KeepTogether(sig_table))

        doc.build(story, canvasmaker=NumberedCanvas)
        buffer.seek(0)
        return buffer.getvalue()

    # --------------------------------------------------------------------------
    # Milestone 6.2: Master Tabulation Register (TR Sheet - Excel Export)
    # --------------------------------------------------------------------------

    @classmethod
    def get_class_tr_data(cls, class_id: int, term_id: int) -> Optional[Dict[str, Any]]:
        """
        Assemble comprehensive tabular data for class Master Tabulation Register.
        """
        target_class = db.session.get(Class, class_id)
        term = db.session.get(ExamTerm, term_id)
        if not target_class or not term:
            return None

        academic_year = target_class.academic_year

        # All mapped subjects for class
        mappings = (
            ClassSubject.query.filter_by(class_id=class_id)
            .join(Subject)
            .order_by(Subject.subject_code.asc())
            .all()
        )

        students = Student.query.filter_by(class_id=class_id).order_by(Student.enrollment_no.asc()).all()

        # Collect marks map: (student_id, class_subject_id) -> Marks
        marks_list = Marks.query.filter_by(term_id=term_id).all()
        marks_by_st_cs = {(m.student_id, m.class_subject_id): m for m in marks_list}

        total_class_max = sum(cs.max_total_marks for cs in mappings)
        student_rows = []

        for st in students:
            sub_scores = []
            st_total = 0.0
            st_has_fail = False
            st_all_absent = True
            st_credits_earned = 0
            st_gp_sum = 0.0

            for cs in mappings:
                m = marks_by_st_cs.get((st.student_id, cs.class_subject_id))
                sub_ref = cs.subject_ref
                credits = sub_ref.credits if sub_ref else 4

                if m:
                    if not m.is_absent:
                        st_all_absent = False
                    in_m = float(m.internal_marks or 0.0)
                    ex_m = float(m.external_marks or 0.0)
                    tot_m = float(m.total_marks or 0.0)
                    grd = m.grade or 'F'
                    gp = float(m.grade_point or 0.0)
                    is_p = bool(m.is_passed)
                    is_ab = bool(m.is_absent)

                    st_total += tot_m
                    st_gp_sum += (gp * credits)
                    if is_p:
                        st_credits_earned += credits
                    else:
                        st_has_fail = True

                    sub_scores.append({
                        'class_subject_id': cs.class_subject_id,
                        'subject_code': sub_ref.subject_code if sub_ref else '',
                        'internal': 'AB' if is_ab else in_m,
                        'external': 'AB' if is_ab else ex_m,
                        'total': 'AB' if is_ab else tot_m,
                        'grade': 'AB' if is_ab else grd,
                        'is_passed': is_p,
                        'is_absent': is_ab
                    })
                else:
                    st_has_fail = True
                    sub_scores.append({
                        'class_subject_id': cs.class_subject_id,
                        'subject_code': sub_ref.subject_code if sub_ref else '',
                        'internal': '-',
                        'external': '-',
                        'total': '-',
                        'grade': '-',
                        'is_passed': False,
                        'is_absent': False
                    })

            st_pct = round((st_total / total_class_max * 100), 2) if total_class_max > 0 else 0.0

            if not mappings:
                status = 'N/A'
            elif st_all_absent and marks_by_st_cs:
                status = 'ABSENT'
            elif st_has_fail:
                status = 'FAIL'
            else:
                status = 'PASS'

            student_rows.append({
                'student_id': st.student_id,
                'enrollment_no': st.enrollment_no,
                'full_name': st.full_name,
                'subject_scores': sub_scores,
                'grand_total': st_total,
                'max_marks': total_class_max,
                'percentage': st_pct,
                'result': status,
                'has_fail': st_has_fail
            })

        # Rank students by grand total
        sorted_students = sorted(student_rows, key=lambda s: s['grand_total'], reverse=True)
        for rank_idx, s in enumerate(sorted_students, start=1):
            s['rank'] = rank_idx if s['grand_total'] > 0 else '-'

        # Cohort summary statistics
        total_enrolled = len(students)
        appeared = sum(1 for s in student_rows if s['result'] != 'ABSENT' and s['grand_total'] > 0)
        passed = sum(1 for s in student_rows if s['result'] == 'PASS')
        failed = sum(1 for s in student_rows if s['result'] == 'FAIL')
        pass_pct = round((passed / appeared * 100), 2) if appeared > 0 else 0.0

        return {
            'class_id': target_class.class_id,
            'class_name': target_class.class_name,
            'section': target_class.section,
            'display_class': target_class.display_name,
            'academic_year': academic_year.year_label if academic_year else 'N/A',
            'term_id': term.term_id,
            'term_name': term.term_name,
            'subjects': [
                {
                    'class_subject_id': cs.class_subject_id,
                    'subject_code': cs.subject_ref.subject_code if cs.subject_ref else '',
                    'subject_name': cs.subject_ref.subject_name if cs.subject_ref else '',
                    'max_internal': cs.max_internal_marks,
                    'max_external': cs.max_external_marks,
                    'max_total': cs.max_total_marks,
                    'pass_marks': cs.pass_marks
                }
                for cs in mappings
            ],
            'students': student_rows,
            'total_class_max': total_class_max,
            'total_enrolled': total_enrolled,
            'total_appeared': appeared,
            'total_passed': passed,
            'total_failed': failed,
            'pass_percentage': pass_pct,
            'export_date': datetime.now(timezone.utc).strftime("%d-%b-%Y %H:%M UTC")
        }

    @classmethod
    def generate_class_tr_sheet_excel(cls, class_id: int, term_id: int) -> Optional[bytes]:
        """
        Generate multi-column Master Tabulation Register Excel (.xlsx) using openpyxl.
        Features merged headers, styled rows, alternating colors, and summary metrics.
        """
        data = cls.get_class_tr_data(class_id, term_id)
        if not data:
            return None

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tabulation Register"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        font_univ = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        font_title = Font(name="Calibri", size=13, bold=True, color="0F172A")
        font_meta = Font(name="Calibri", size=10, italic=True, color="475569")
        font_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_subhdr = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        font_cell = Font(name="Calibri", size=9, color="0F172A")
        font_cell_bold = Font(name="Calibri", size=9, bold=True, color="0F172A")
        font_pass = Font(name="Calibri", size=9, bold=True, color="15803D")
        font_fail = Font(name="Calibri", size=9, bold=True, color="B91C1C")

        fill_hdr_primary = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy
        fill_hdr_secondary = PatternFill(start_color="334155", end_color="334155", fill_type="solid")  # Slate
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        fill_fail = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        fill_total_row = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        thin_side = Side(border_style="thin", color="CBD5E1")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Side(border_style="double", color="1E293B")
        border_summary = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)

        num_subjects = len(data['subjects'])
        # Fixed cols: Sl No (1), Enrollment (2), Name (3)
        # Subjects cols: 4 per subject (Int, Ext, Tot, Grd)
        # Summary cols: Grand Total, Max Marks, Pct, Result, Rank (5 cols)
        total_cols = 3 + (num_subjects * 4) + 5

        # 1. Title Banner Rows
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        c1 = ws.cell(row=1, column=1, value="INDIRA GANDHI NATIONAL OPEN UNIVERSITY")
        c1.font = font_univ
        c1.alignment = align_center

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
        c2 = ws.cell(row=2, column=1, value="MASTER TABULATION REGISTER (TR SHEET)")
        c2.font = font_title
        c2.alignment = align_center

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
        meta_str = (
            f"Class: {data['display_class']} | Session: {data['academic_year']} | "
            f"Examination Term: {data['term_name']} | Generated: {data['export_date']}"
        )
        c3 = ws.cell(row=3, column=1, value=meta_str)
        c3.font = font_meta
        c3.alignment = align_center

        ws.row_dimensions[1].height = 25
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[4].height = 8

        # 2. Table Headers (Row 5 & 6)
        hdr_row1 = 5
        hdr_row2 = 6
        ws.row_dimensions[hdr_row1].height = 24
        ws.row_dimensions[hdr_row2].height = 20

        # Fixed Metadata Columns
        fixed_hdrs = [
            ("Sl.", 1),
            ("Enrollment No", 2),
            ("Student Name", 3)
        ]
        for title, col_idx in fixed_hdrs:
            ws.merge_cells(start_row=hdr_row1, start_column=col_idx, end_row=hdr_row2, end_column=col_idx)
            cell = ws.cell(row=hdr_row1, column=col_idx, value=title)
            cell.font = font_hdr
            cell.fill = fill_hdr_primary
            cell.alignment = align_center
            cell.border = border_all
            ws.cell(row=hdr_row2, column=col_idx).border = border_all

        cur_col = 4
        # Subject Columns
        for sub in data['subjects']:
            start_c = cur_col
            end_c = cur_col + 3
            ws.merge_cells(start_row=hdr_row1, start_column=start_c, end_row=hdr_row1, end_column=end_c)
            sub_title = f"{sub['subject_code']} (Max {int(sub['max_total'])})"
            cell = ws.cell(row=hdr_row1, column=start_c, value=sub_title)
            cell.font = font_hdr
            cell.fill = fill_hdr_primary
            cell.alignment = align_center
            for c_idx in range(start_c, end_c + 1):
                ws.cell(row=hdr_row1, column=c_idx).border = border_all

            # Sub-headers (Row 6)
            sub_hdrs = [
                f"Int ({int(sub['max_internal'])})",
                f"Ext ({int(sub['max_external'])})",
                "Tot",
                "Grd"
            ]
            for offset, sh in enumerate(sub_hdrs):
                sc = ws.cell(row=hdr_row2, column=start_c + offset, value=sh)
                sc.font = font_subhdr
                sc.fill = fill_hdr_secondary
                sc.alignment = align_center
                sc.border = border_all

            cur_col += 4

        # Trailing Summary Columns
        summary_hdrs = [
            ("Grand Total", cur_col),
            ("Max Marks", cur_col + 1),
            ("Percentage", cur_col + 2),
            ("Result", cur_col + 3),
            ("Rank", cur_col + 4)
        ]
        for title, col_idx in summary_hdrs:
            ws.merge_cells(start_row=hdr_row1, start_column=col_idx, end_row=hdr_row2, end_column=col_idx)
            cell = ws.cell(row=hdr_row1, column=col_idx, value=title)
            cell.font = font_hdr
            cell.fill = fill_hdr_primary
            cell.alignment = align_center
            cell.border = border_all
            ws.cell(row=hdr_row2, column=col_idx).border = border_all

        # 3. Data Rows
        current_data_row = 7
        for idx, st in enumerate(data['students'], start=1):
            is_even = (idx % 2 == 0)
            row_fill = fill_zebra if is_even else None
            ws.row_dimensions[current_data_row].height = 18

            # Fixed metadata
            c_sl = ws.cell(row=current_data_row, column=1, value=idx)
            c_sl.alignment = align_center
            c_enr = ws.cell(row=current_data_row, column=2, value=st['enrollment_no'])
            c_enr.alignment = align_center
            c_name = ws.cell(row=current_data_row, column=3, value=st['full_name'])
            c_name.alignment = align_left

            col_pointer = 4
            for sc in st['subject_scores']:
                c_in = ws.cell(row=current_data_row, column=col_pointer, value=sc['internal'])
                c_in.alignment = align_center
                c_ex = ws.cell(row=current_data_row, column=col_pointer + 1, value=sc['external'])
                c_ex.alignment = align_center
                c_tot = ws.cell(row=current_data_row, column=col_pointer + 2, value=sc['total'])
                c_tot.alignment = align_center
                c_grd = ws.cell(row=current_data_row, column=col_pointer + 3, value=sc['grade'])
                c_grd.alignment = align_center
                col_pointer += 4

            # Summary cols
            c_gtot = ws.cell(row=current_data_row, column=col_pointer, value=st['grand_total'])
            c_gtot.alignment = align_center
            c_max = ws.cell(row=current_data_row, column=col_pointer + 1, value=st['max_marks'])
            c_max.alignment = align_center
            c_pct = ws.cell(row=current_data_row, column=col_pointer + 2, value=f"{st['percentage']:.2f}%")
            c_pct.alignment = align_center

            c_res = ws.cell(row=current_data_row, column=col_pointer + 3, value=st['result'])
            c_res.alignment = align_center
            if st['result'] == 'PASS':
                c_res.font = font_pass
                c_res.fill = fill_pass
            elif st['result'] in ['FAIL', 'REAPPEAR']:
                c_res.font = font_fail
                c_res.fill = fill_fail

            c_rnk = ws.cell(row=current_data_row, column=col_pointer + 4, value=st['rank'])
            c_rnk.alignment = align_center

            # Apply borders and fonts
            for col_i in range(1, total_cols + 1):
                cell = ws.cell(row=current_data_row, column=col_i)
                if not cell.font or cell.font == font_cell:
                    cell.font = font_cell
                cell.border = border_all
                if row_fill and cell != c_res:
                    cell.fill = row_fill

            current_data_row += 1

        # 4. Cohort Summary Rows at Bottom
        ws.row_dimensions[current_data_row].height = 20
        ws.merge_cells(start_row=current_data_row, start_column=1, end_row=current_data_row, end_column=3)
        lbl_cell = ws.cell(row=current_data_row, column=1, value="COHORT SUMMARY")
        lbl_cell.font = font_cell_bold
        lbl_cell.alignment = align_center
        lbl_cell.fill = fill_total_row

        summary_note = (
            f"Enrolled: {data['total_enrolled']} | Appeared: {data['total_appeared']} | "
            f"Passed: {data['total_passed']} | Failed: {data['total_failed']} | "
            f"Pass Rate: {data['pass_percentage']:.1f}%"
        )
        ws.merge_cells(start_row=current_data_row, start_column=4, end_row=current_data_row, end_column=total_cols)
        note_cell = ws.cell(row=current_data_row, column=4, value=summary_note)
        note_cell.font = font_cell_bold
        note_cell.alignment = align_center
        note_cell.fill = fill_total_row

        for col_i in range(1, total_cols + 1):
            ws.cell(row=current_data_row, column=col_i).border = border_summary

        # 5. Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # skip title banner rows for width calculation
                if cell.row in [1, 2, 3, current_data_row]:
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 10)

        # Ensure Sl No and Name columns have generous widths
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 24

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    # --------------------------------------------------------------------------
    # Milestone 6.3: Executive Class Performance Summary Report
    # --------------------------------------------------------------------------

    @classmethod
    def get_class_summary_data(cls, class_id: int, term_id: int) -> Optional[Dict[str, Any]]:
        """
        Assemble executive analytics and performance metrics for Faculty/HOD summary.
        """
        target_class = db.session.get(Class, class_id)
        term = db.session.get(ExamTerm, term_id)
        if not target_class or not term:
            return None

        # Gather class subject mappings for metadata
        mappings = {
            cs.subject_ref.subject_code: cs
            for cs in ClassSubject.query.filter_by(class_id=class_id).all()
            if cs.subject_ref
        }

        overview = AnalyticsService.compute_overview_metrics(class_id, term_id)
        overview['mean_score'] = overview.get('class_average', 0.0)
        subj_comp = AnalyticsService.compute_subject_comparison(class_id, term_id)
        raw_subjects = subj_comp.get('subjects', [])

        enriched_subjects = []
        for s in raw_subjects:
            code = s.get('subject_code')
            mapping = mappings.get(code)
            app = s.get('appeared', 0)
            psd = s.get('passed', 0)
            fld = s.get('failed', 0)
            prate = s.get('pass_rate', 0.0)
            avg_m = s.get('average_marks', 0.0)
            hi_m = s.get('highest_marks', 0.0)
            lo_m = s.get('lowest_marks', 0.0)

            enriched_subjects.append({
                'subject_id': s.get('subject_id'),
                'subject_code': code,
                'subject_name': s.get('subject_name'),
                'max_total': mapping.max_total_marks if mapping else 100.0,
                'pass_marks': mapping.pass_marks if mapping else 40.0,
                'total_candidates': s.get('total_candidates', 0),
                'appeared': app,
                'total_appeared': app,
                'absent': s.get('absent', 0),
                'passed': psd,
                'passed_count': psd,
                'failed': fld,
                'failed_count': fld,
                'pass_rate': prate,
                'pass_percentage': prate,
                'average_marks': avg_m,
                'average_score': avg_m,
                'average_percentage': s.get('average_percentage', 0.0),
                'highest_marks': hi_m,
                'highest_score': hi_m,
                'lowest_marks': lo_m,
                'lowest_score': lo_m
            })

        grade_dist = AnalyticsService.compute_grade_distribution(class_id, term_id)
        at_risk = AnalyticsService.identify_at_risk_students(class_id, term_id)
        top_performers = AnalyticsService.compute_top_performers(class_id, term_id, limit=5)

        return {
            'class_id': target_class.class_id,
            'class_name': target_class.class_name,
            'section': target_class.section,
            'display_class': target_class.display_name,
            'academic_year': target_class.academic_year.year_label if target_class.academic_year else 'N/A',
            'term_id': term.term_id,
            'term_name': term.term_name,
            'overview': overview,
            'subjects': enriched_subjects,
            'subj_comp': subj_comp,
            'grade_dist': grade_dist,
            'at_risk': at_risk,
            'top_performers': top_performers,
            'generated_at': datetime.now(timezone.utc).strftime("%d-%b-%Y %H:%M UTC")
        }

    @classmethod
    def generate_class_summary_pdf(cls, class_id: int, term_id: int) -> Optional[bytes]:
        """
        Generate executive printable PDF performance summary for HOD / Academic Council.
        """
        data = cls.get_class_summary_data(class_id, term_id)
        if not data:
            return None

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'SumTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            alignment=1,
            textColor=colors.HexColor('#0f172a')
        )
        sub_title_style = ParagraphStyle(
            'SumSubTitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            alignment=1,
            textColor=colors.HexColor('#475569')
        )
        section_hdr = ParagraphStyle(
            'SecHdr',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#1e3a8a')
        )
        tbl_hdr = ParagraphStyle(
            'TblHdr',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=10,
            alignment=1,
            textColor=colors.white
        )
        tbl_cell = ParagraphStyle(
            'TblCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            alignment=0,
            textColor=colors.HexColor('#1e293b')
        )
        tbl_cell_center = ParagraphStyle(
            'TblCellCenter',
            parent=tbl_cell,
            alignment=1
        )
        tbl_cell_bold = ParagraphStyle(
            'TblCellBold',
            parent=tbl_cell,
            fontName='Helvetica-Bold'
        )

        story = []

        # 1. Header
        story.append(Paragraph("INDIRA GANDHI NATIONAL OPEN UNIVERSITY", title_style))
        story.append(Paragraph("EXECUTIVE CLASS PERFORMANCE & STATISTICAL RESULT SUMMARY", ParagraphStyle('ExSub', parent=title_style, fontSize=12, leading=15, textColor=colors.HexColor('#1e3a8a'))))
        story.append(Spacer(1, 2))
        meta_str = f"Class: <b>{data['display_class']}</b> | Academic Session: <b>{data['academic_year']}</b> | Examination Term: <b>{data['term_name']}</b> | Report Date: {data['generated_at']}"
        story.append(Paragraph(meta_str, sub_title_style))
        story.append(Spacer(1, 4))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1e3a8a'), spaceBefore=2, spaceAfter=8))

        # 2. Executive KPI Overview Cards
        ov = data['overview']
        kpi_data = [
            [
                Paragraph("<b>Total Enrolled:</b>", tbl_cell_bold), Paragraph(str(ov.get('total_enrolled', 0)), tbl_cell_center),
                Paragraph("<b>Total Appeared:</b>", tbl_cell_bold), Paragraph(str(ov.get('total_appeared', 0)), tbl_cell_center),
                Paragraph("<b>Total Passed:</b>", tbl_cell_bold), Paragraph(f"<font color='#15803d'><b>{ov.get('total_passed', 0)}</b></font>", tbl_cell_center),
                Paragraph("<b>Total Failed:</b>", tbl_cell_bold), Paragraph(f"<font color='#b91c1c'><b>{ov.get('total_failed', 0)}</b></font>", tbl_cell_center),
            ],
            [
                Paragraph("<b>Pass Percentage:</b>", tbl_cell_bold), Paragraph(f"<b>{ov.get('pass_percentage', 0.0):.1f}%</b>", tbl_cell_center),
                Paragraph("<b>Class Average:</b>", tbl_cell_bold), Paragraph(f"<b>{ov.get('mean_score', 0.0):.2f}</b>", tbl_cell_center),
                Paragraph("<b>Median Score:</b>", tbl_cell_bold), Paragraph(f"{ov.get('median_score', 0.0):.2f}", tbl_cell_center),
                Paragraph("<b>Highest / Lowest:</b>", tbl_cell_bold), Paragraph(f"{ov.get('highest_score', 0.0):.1f} / {ov.get('lowest_score', 0.0):.1f}", tbl_cell_center),
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[100, 90, 100, 90, 100, 90, 100, 90])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#94a3b8')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 12))

        # 3. Subject-wise Comparison Table
        story.append(Paragraph("Subject-Wise Academic Performance Analysis", section_hdr))
        story.append(Spacer(1, 4))

        sub_headers = [
            Paragraph("Code", tbl_hdr),
            Paragraph("Subject Title", tbl_hdr),
            Paragraph("Max Marks", tbl_hdr),
            Paragraph("Pass Marks", tbl_hdr),
            Paragraph("Appeared", tbl_hdr),
            Paragraph("Passed", tbl_hdr),
            Paragraph("Failed", tbl_hdr),
            Paragraph("Pass %", tbl_hdr),
            Paragraph("Avg Score", tbl_hdr),
            Paragraph("Top Score", tbl_hdr),
        ]
        sub_rows = [sub_headers]
        for sb in data['subjects']:
            pass_rate = sb.get('pass_rate', 0.0)
            pass_color = '#15803d' if pass_rate >= 50.0 else '#b91c1c'
            sub_rows.append([
                Paragraph(sb.get('subject_code', ''), tbl_cell_bold),
                Paragraph(sb.get('subject_name', ''), tbl_cell),
                Paragraph(str(int(sb.get('max_total', 100))), tbl_cell_center),
                Paragraph(str(int(sb.get('pass_marks', 40))), tbl_cell_center),
                Paragraph(str(sb.get('appeared', 0)), tbl_cell_center),
                Paragraph(str(sb.get('passed', 0)), tbl_cell_center),
                Paragraph(str(sb.get('failed', 0)), tbl_cell_center),
                Paragraph(f"<font color='{pass_color}'><b>{pass_rate:.1f}%</b></font>", tbl_cell_center),
                Paragraph(f"{sb.get('average_marks', 0.0):.1f}", tbl_cell_center),
                Paragraph(f"{sb.get('highest_marks', 0.0):.1f}", tbl_cell_center),
            ])

        sub_table = Table(sub_rows, colWidths=[65, 235, 55, 55, 55, 50, 50, 65, 65, 65])
        sub_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#1e3a8a')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))
        story.append(sub_table)
        story.append(Spacer(1, 12))

        # 4. Grade Distribution & Academic Alert Candidates (Side-by-side)
        story.append(Paragraph("Grade Distribution & At-Risk Candidates", section_hdr))
        story.append(Spacer(1, 4))

        # Grade Distribution table
        gd = data['grade_dist']
        gd_labels = gd.get('labels', [])
        gd_counts = gd.get('counts', [])
        gd_pcts = gd.get('percentages', [])

        gd_rows = [[
            Paragraph("Grade", tbl_hdr),
            Paragraph("Count", tbl_hdr),
            Paragraph("Percentage", tbl_hdr),
        ]]
        for lbl, cnt, pct in zip(gd_labels, gd_counts, gd_pcts):
            gd_rows.append([
                Paragraph(f"<b>{lbl}</b>", tbl_cell_center),
                Paragraph(str(cnt), tbl_cell_center),
                Paragraph(f"{pct:.1f}%", tbl_cell_center),
            ])

        gd_table = Table(gd_rows, colWidths=[60, 60, 80])
        gd_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#334155')),
            ('TOPPADDING', (0, 0), (-1, -1), 2.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))

        # At-Risk Table
        ar_list = data['at_risk'][:5]  # Top 5 at risk
        ar_rows = [[
            Paragraph("Enrollment", tbl_hdr),
            Paragraph("Student Name", tbl_hdr),
            Paragraph("Risk Severity", tbl_hdr),
            Paragraph("Recommended Action", tbl_hdr),
        ]]
        if not ar_list:
            ar_rows.append([
                Paragraph("-", tbl_cell_center),
                Paragraph("No at-risk candidates identified.", tbl_cell),
                Paragraph("-", tbl_cell_center),
                Paragraph("Satisfactory academic progress.", tbl_cell)
            ])
        else:
            for ar in ar_list:
                sev = ar.get('risk_severity', 'MEDIUM')
                sev_color = '#b91c1c' if sev == 'HIGH' else '#f59e0b'
                ar_rows.append([
                    Paragraph(ar.get('enrollment_no', ''), tbl_cell_bold),
                    Paragraph(ar.get('full_name', ''), tbl_cell),
                    Paragraph(f"<font color='{sev_color}'><b>{sev}</b></font>", tbl_cell_center),
                    Paragraph(ar.get('recommended_action', 'Academic counseling'), tbl_cell),
                ])

        ar_table = Table(ar_rows, colWidths=[90, 150, 80, 230])
        ar_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#334155')),
            ('TOPPADDING', (0, 0), (-1, -1), 2.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))

        combo_table = Table([[gd_table, ar_table]], colWidths=[210, 550])
        combo_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(combo_table)
        story.append(Spacer(1, 16))

        # Signatures
        sig_data = [
            [
                Paragraph("__________________________<br/><b>Faculty Coordinator / Program In-Charge</b>", ParagraphStyle('SigF', parent=styles['Normal'], alignment=0)),
                Paragraph("__________________________<br/><b>Head of Department / Dean</b>", ParagraphStyle('SigH', parent=styles['Normal'], alignment=2)),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[380, 380])
        sig_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'BOTTOM'),
        ]))
        story.append(KeepTogether(sig_table))

        doc.build(story, canvasmaker=NumberedCanvas)
        buffer.seek(0)
        return buffer.getvalue()
