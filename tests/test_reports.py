"""Unit and Integration Tests for Phase 6 Reporting, Document Generation & Export Services."""

from io import BytesIO
import pytest
import openpyxl
from app import create_app, db
from app.models import (
    User, AcademicYear, Class, Subject, ClassSubject, Student,
    ExamTerm, Marks, GradingScale, GradeRule
)
from app.services.report_service import ReportService


@pytest.fixture
def app():
    """Create test application with in-memory SQLite and seeded test data."""
    test_app = create_app('testing')

    with test_app.app_context():
        db.create_all()

        # 1. Users
        admin = User(username='admin', full_name='Admin User', role=User.ROLE_ADMIN)
        admin.set_password('Admin@123')

        teacher = User(username='teacher', full_name='Teacher User', role=User.ROLE_TEACHER)
        teacher.set_password('Teacher@123')

        student_user = User(username='student', full_name='Student User', role=User.ROLE_STUDENT)
        student_user.set_password('Student@123')

        db.session.add_all([admin, teacher, student_user])

        # 2. Grading Scale
        scale = GradingScale(scale_name="IGNOU Standard 10-Point", is_active=True)
        db.session.add(scale)
        db.session.flush()

        rules_data = [
            (85.0, 100.0, 'O', 10.0, 'Outstanding'),
            (75.0, 84.99, 'A+', 9.0, 'Excellent'),
            (65.0, 74.99, 'A', 8.0, 'Very Good'),
            (55.0, 64.99, 'B+', 7.0, 'Good'),
            (50.0, 54.99, 'B', 6.0, 'Above Average'),
            (40.0, 49.99, 'C', 5.0, 'Average'),
            (0.0, 39.99, 'F', 0.0, 'Fail')
        ]
        for r in rules_data:
            db.session.add(GradeRule(
                scale_id=scale.scale_id,
                min_percentage=r[0],
                max_percentage=r[1],
                grade_letter=r[2],
                grade_point=r[3],
                description=r[4]
            ))

        # 3. Academic Structure
        year = AcademicYear(year_label='2025-2026', is_active=True)
        db.session.add(year)
        db.session.flush()

        cls1 = Class(class_name='BCA-Sem-1', section='A', year_id=year.year_id)
        term1 = ExamTerm(term_name='TEE June 2025', year_id=year.year_id, is_locked=False)
        sub1 = Subject(subject_code='BCS-011', subject_name='Computer Basics', credits=3)
        sub2 = Subject(subject_code='BCS-012', subject_name='Mathematics', credits=4)
        db.session.add_all([cls1, term1, sub1, sub2])
        db.session.flush()

        cs1 = ClassSubject(class_id=cls1.class_id, subject_id=sub1.subject_id, max_internal_marks=30.0, max_external_marks=70.0, pass_marks=40.0)
        cs2 = ClassSubject(class_id=cls1.class_id, subject_id=sub2.subject_id, max_internal_marks=30.0, max_external_marks=70.0, pass_marks=40.0)
        db.session.add_all([cs1, cs2])
        db.session.flush()

        # 4. Students
        s1 = Student(enrollment_no='240100101', full_name='Aarav Sharma', class_id=cls1.class_id)
        s2 = Student(enrollment_no='240100102', full_name='Diya Patel', class_id=cls1.class_id)
        s3 = Student(enrollment_no='240100103', full_name='Kabir Verma', class_id=cls1.class_id)
        db.session.add_all([s1, s2, s3])
        db.session.flush()

        # 5. Marks
        # s1: passed both courses
        m1_1 = Marks(student_id=s1.student_id, class_subject_id=cs1.class_subject_id, term_id=term1.term_id, internal_marks=25.0, external_marks=60.0, total_marks=85.0, percentage=85.0, grade='O', grade_point=10.0, is_passed=True)
        m1_2 = Marks(student_id=s1.student_id, class_subject_id=cs2.class_subject_id, term_id=term1.term_id, internal_marks=22.0, external_marks=56.0, total_marks=78.0, percentage=78.0, grade='A+', grade_point=9.0, is_passed=True)

        # s2: passed one, failed one
        m2_1 = Marks(student_id=s2.student_id, class_subject_id=cs1.class_subject_id, term_id=term1.term_id, internal_marks=18.0, external_marks=42.0, total_marks=60.0, percentage=60.0, grade='B+', grade_point=7.0, is_passed=True)
        m2_2 = Marks(student_id=s2.student_id, class_subject_id=cs2.class_subject_id, term_id=term1.term_id, internal_marks=10.0, external_marks=20.0, total_marks=30.0, percentage=30.0, grade='F', grade_point=0.0, is_passed=False)

        # s3: absent
        m3_1 = Marks(student_id=s3.student_id, class_subject_id=cs1.class_subject_id, term_id=term1.term_id, internal_marks=0.0, external_marks=0.0, total_marks=0.0, percentage=0.0, grade='F', grade_point=0.0, is_absent=True, is_passed=False)
        m3_2 = Marks(student_id=s3.student_id, class_subject_id=cs2.class_subject_id, term_id=term1.term_id, internal_marks=0.0, external_marks=0.0, total_marks=0.0, percentage=0.0, grade='F', grade_point=0.0, is_absent=True, is_passed=False)

        db.session.add_all([m1_1, m1_2, m2_1, m2_2, m3_1, m3_2])
        db.session.commit()

    yield test_app


@pytest.fixture
def client(app):
    """Test HTTP client."""
    return app.test_client()


def login(client, username, password):
    """Helper to authenticate user session."""
    return client.post('/login', data={'username': username, 'password': password}, follow_redirects=True)


# ==============================================================================
# Milestone 6.1 Tests: Student Report Card
# ==============================================================================

def test_get_student_report_card_data(app):
    """Verify aggregation of student marksheet data."""
    with app.app_context():
        s1 = Student.query.filter_by(enrollment_no='240100101').first()
        term = ExamTerm.query.first()

        data = ReportService.get_student_report_card_data(s1.student_id, term.term_id)
        assert data is not None
        assert data['enrollment_no'] == '240100101'
        assert data['full_name'] == 'Aarav Sharma'
        assert len(data['subjects']) == 2
        assert data['total_max_marks'] == 200.0
        assert data['total_obtained_marks'] == 163.0
        assert data['overall_percentage'] == 81.5
        assert data['has_fail'] is False
        assert "PASSED" in data['overall_result']
        assert data['class_rank'] == 1


def test_student_report_card_with_failure(app):
    """Verify student marksheet data when one subject failed."""
    with app.app_context():
        s2 = Student.query.filter_by(enrollment_no='240100102').first()
        term = ExamTerm.query.first()

        data = ReportService.get_student_report_card_data(s2.student_id, term.term_id)
        assert data is not None
        assert data['has_fail'] is True
        assert data['overall_result'] == "REAPPEAR / FAILED"


def test_generate_student_report_card_pdf(app):
    """Verify ReportLab PDF generation generates valid non-empty PDF bytes."""
    with app.app_context():
        s1 = Student.query.filter_by(enrollment_no='240100101').first()
        term = ExamTerm.query.first()

        pdf_bytes = ReportService.generate_student_report_card_pdf(s1.student_id, term.term_id)
        assert pdf_bytes is not None
        assert isinstance(pdf_bytes, bytes)
        assert len(pdf_bytes) > 1000
        # Standard PDF magic bytes
        assert pdf_bytes.startswith(b'%PDF-')


def test_student_report_card_not_found(app):
    """Verify invalid student/term IDs gracefully return None."""
    with app.app_context():
        assert ReportService.get_student_report_card_data(9999, 1) is None
        assert ReportService.generate_student_report_card_pdf(9999, 1) is None


# ==============================================================================
# Milestone 6.2 Tests: Master Tabulation Register (TR Sheet)
# ==============================================================================

def test_get_class_tr_data(app):
    """Verify comprehensive class tabulation data assembly."""
    with app.app_context():
        cls1 = Class.query.first()
        term = ExamTerm.query.first()

        tr_data = ReportService.get_class_tr_data(cls1.class_id, term.term_id)
        assert tr_data is not None
        assert tr_data['class_id'] == cls1.class_id
        assert len(tr_data['subjects']) == 2
        assert tr_data['total_enrolled'] == 3
        assert tr_data['total_appeared'] == 2  # s1 and s2 appeared, s3 absent
        assert tr_data['total_passed'] == 1    # s1 passed
        assert tr_data['total_failed'] == 1    # s2 failed
        assert tr_data['pass_percentage'] == 50.0

        # Verify student records
        students = tr_data['students']
        assert len(students) == 3
        s1_row = next(s for s in students if s['enrollment_no'] == '240100101')
        assert s1_row['grand_total'] == 163.0
        assert s1_row['result'] == 'PASS'
        assert s1_row['rank'] == 1


def test_generate_class_tr_sheet_excel(app):
    """Verify openpyxl generation of Master Tabulation Register Excel workbook."""
    with app.app_context():
        cls1 = Class.query.first()
        term = ExamTerm.query.first()

        excel_bytes = ReportService.generate_class_tr_sheet_excel(cls1.class_id, term.term_id)
        assert excel_bytes is not None
        assert len(excel_bytes) > 1000

        # Load generated workbook with openpyxl and inspect structure
        wb = openpyxl.load_workbook(BytesIO(excel_bytes))
        assert "Tabulation Register" in wb.sheetnames
        ws = wb["Tabulation Register"]

        # Check university banner
        assert "INDIRA GANDHI NATIONAL OPEN UNIVERSITY" in str(ws.cell(row=1, column=1).value)
        assert "MASTER TABULATION REGISTER" in str(ws.cell(row=2, column=1).value)

        # Check metadata row
        assert "BCA-Sem-1" in str(ws.cell(row=3, column=1).value)

        # Check header titles
        assert ws.cell(row=5, column=1).value == "Sl."
        assert ws.cell(row=5, column=2).value == "Enrollment No"
        assert ws.cell(row=5, column=3).value == "Student Name"

        # Check student data row 7 (first student)
        assert ws.cell(row=7, column=2).value == '240100101'
        assert ws.cell(row=7, column=3).value == 'Aarav Sharma'


# ==============================================================================
# Milestone 6.3 Tests: Executive Class Performance Summary
# ==============================================================================

def test_get_class_summary_data(app):
    """Verify executive class summary metrics integration."""
    with app.app_context():
        cls1 = Class.query.first()
        term = ExamTerm.query.first()

        sum_data = ReportService.get_class_summary_data(cls1.class_id, term.term_id)
        assert sum_data is not None
        assert 'overview' in sum_data
        assert 'subjects' in sum_data
        assert 'grade_dist' in sum_data
        assert 'at_risk' in sum_data
        assert sum_data['overview']['total_enrolled'] == 3


def test_generate_class_summary_pdf(app):
    """Verify ReportLab generation of executive performance summary landscape PDF."""
    with app.app_context():
        cls1 = Class.query.first()
        term = ExamTerm.query.first()

        pdf_bytes = ReportService.generate_class_summary_pdf(cls1.class_id, term.term_id)
        assert pdf_bytes is not None
        assert isinstance(pdf_bytes, bytes)
        assert len(pdf_bytes) > 1000
        assert pdf_bytes.startswith(b'%PDF-')


# ==============================================================================
# Route & Access Control Tests
# ==============================================================================

def test_reports_access_control(client):
    """Ensure report endpoints are protected by teacher_required."""
    resp = client.get('/reports/report-cards')
    assert resp.status_code == 302
    assert '/login' in resp.headers['Location']

    resp = client.get('/reports/tabulation-register')
    assert resp.status_code == 302

    resp = client.get('/reports/class-summary')
    assert resp.status_code == 302


def test_reports_pages_authenticated(client, app):
    """Verify authenticated teacher can access report pages and endpoints."""
    login(client, 'teacher', 'Teacher@123')

    with app.app_context():
        cls1 = Class.query.first()
        term = ExamTerm.query.first()
        s1 = Student.query.first()
        cid = cls1.class_id
        tid = term.term_id
        sid = s1.student_id

    # 1. Report Cards Roster
    res = client.get(f'/reports/report-cards?class_id={cid}&term_id={tid}')
    assert res.status_code == 200
    assert b"Institutional Student Report Cards" in res.data
    assert b"Aarav Sharma" in res.data

    # 2. Report Card Preview
    res = client.get(f'/reports/report-card/{sid}?term_id={tid}')
    assert res.status_code == 200
    assert b"Official Grade Card" in res.data
    assert b"240100101" in res.data

    # 3. Report Card PDF Download
    res = client.get(f'/reports/report-card/{sid}/pdf?term_id={tid}')
    assert res.status_code == 200
    assert res.mimetype == 'application/pdf'
    assert b'%PDF-' in res.data

    # 4. Tabulation Register Web View
    res = client.get(f'/reports/tabulation-register?class_id={cid}&term_id={tid}')
    assert res.status_code == 200
    assert b"Master Tabulation Register" in res.data
    assert b"BCS-011" in res.data

    # 5. Tabulation Register Excel Export
    res = client.get(f'/reports/tabulation-register/export?class_id={cid}&term_id={tid}')
    assert res.status_code == 200
    assert res.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment;' in res.headers['Content-Disposition']

    # 6. Class Summary Web View
    res = client.get(f'/reports/class-summary?class_id={cid}&term_id={tid}')
    assert res.status_code == 200
    assert b"Executive Class Performance Summary" in res.data

    # 7. Class Summary PDF Download
    res = client.get(f'/reports/class-summary/pdf?class_id={cid}&term_id={tid}')
    assert res.status_code == 200
    assert res.mimetype == 'application/pdf'
    assert b'%PDF-' in res.data
